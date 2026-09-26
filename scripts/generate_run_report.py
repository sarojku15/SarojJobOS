#!/usr/bin/env python3

"""
Phase 7 reporting layer: the actionable, multi-sheet daily job-search
workbook. Reuses every existing pure pipeline component -- nothing here
re-implements scoring, eligibility, freshness, or dedup:

    jobs table (source, company, title, ..., score, priority, status)
        -> job_ranking.build_ranking_record()   (existing, unmodified --
           itself reuses job_eligibility.assess_job_eligibility(),
           score_job.score_job(), score_explanation.build_score_explanation(),
           freshness.classify_freshness(), canonical_job.derive_canonical_job())
        -> cross_source_dedup.find_cross_source_duplicate_candidates()
           (existing, unmodified -- run once over the whole batch)
        -> ReportRow (this module's only new data shape)
        -> sheet classification (this module's only new logic)
        -> .xlsx workbook (openpyxl)

This module writes NOTHING to any database -- it only reads. It never
recomputes a score or freshness category differently than the existing
pipeline already does; every score/priority/freshness value shown is
produced by calling the exact same functions search_worker.py itself
calls, not a second, competing implementation.

============================================================================
KNOWN, DOCUMENTED DATA-MODEL LIMITATIONS THIS MODULE WORKS AROUND
============================================================================
(See data/reports/phase7_reporting_audit.md for the full evidence trail.)

1. jobs.discovered_at is UNRELIABLE as "first discovered date" --
   tracker.py's upsert ON CONFLICT clause overwrites it on every
   re-ingest. jobs.created_at (SQLite's own insert-time default,
   never touched by any UPDATE SET clause in this codebase) is used
   instead -- it is the field that actually behaves as "first
   discovered."

2. jobs.freshness is a real column in the schema (added by
   migrate_v2_schema.py) but is NEVER written by any pipeline code.
   This module never reads it -- freshness is always RECOMPUTED via
   freshness.classify_freshness(posted_date), exactly as
   job_ranking.build_ranking_record() already does for search_worker.py's
   own in-memory ranking enrichment.

3. jobs.score/priority/status are effectively single-candidate-era
   columns (see search_worker.py's own docstring on this) -- this
   module re-derives eligibility/score/priority per-candidate, fresh,
   via build_ranking_record(), rather than trusting the possibly
   stale/other-candidate's values in those columns directly. The
   `status` column (application lifecycle) is the one column this
   module DOES read directly from the jobs table, since it is the
   only place that state exists.

4. status="REJECTED" WAS ambiguous in the codebase Phase 7.1 audited:
   score_job() wrote status="REJECTED" for any eligible-but-<70-scored
   job at INSERT time (tracker.upsert_job() writes `status` only in
   the INSERT column list, not in the ON CONFLICT UPDATE SET clause --
   so this initial value persists until something else changes it),
   while application_schema.json's lifecycle also defined "REJECTED"
   to mean an employer rejected a real application -- and no code
   anywhere ever transitioned a job's status through
   FOUND -> ... -> APPLIED -> ... -> REJECTED (no application/approval
   workflow existed, confirmed by tracing the code).

   **FIXED in Phase 7.2** (data/reports/phase7_2_automation_status_audit.md):
   score_job.py now writes "NOT_QUALIFIED" for this outcome, never
   "REJECTED" -- a one-line-per-branch vocabulary change (Phase 7.1's
   recommended Option A), leaving "REJECTED" reserved, going forward,
   for a real employer-rejection value once an application workflow
   exists to write "EMPLOYER_REJECTED" (a new, reserved,
   currently-unwritten status).

   **Backward compatibility, not a migration**: existing rows written
   before this change may still literally say status="REJECTED"
   (confirmed: 3 such rows exist in production today, all
   source='TEST' -- verified by a read-only query, never modified).
   `SCORE_BUCKET_REJECTION_STATUSES = ("NOT_QUALIFIED", "REJECTED")`
   is treated identically everywhere in this module, forever -- an old
   "REJECTED" row is NEVER reinterpreted as an employer rejection, and
   NEVER will be, since this codebase has no mechanism to distinguish
   a legacy score-rejection from a hypothetical old employer-rejection
   that also happened to reuse the same string (there is no evidence
   either way, so the conservative, documented choice --
   `SCORE_REJECTED_NOT_QUALIFIED` -- is always taken, per explicit
   instruction: "do not invent employer rejection history").

============================================================================
PHASE 7.1 ADDITIONS (data/reports/phase7_1_reporting_semantics_audit.md)
============================================================================
5. NEW_JOBS now means "first time THIS CANDIDATE matched this job",
   not merely "first time this job row existed in the global `jobs`
   table". The correct, already-existing, per-candidate signal is
   `candidate_job_matches.created_at` -- PRIMARY KEY (candidate_id,
   job_id), and (confirmed by reading search_worker.py's
   upsert_candidate_job_match()) never touched by its own ON CONFLICT
   UPDATE SET clause, so it behaves exactly like jobs.created_at's
   protection, but per-candidate. Since `candidate_job_matches` rows
   only ever exist for jobs that passed eligibility (search_worker.py
   never writes one for an ineligible job), this signal is available
   for exactly the jobs NEW_JOBS could ever contain (qualifying jobs) --
   no gap. Phase 7's original implementation used the WRONG, global
   `jobs.created_at` instead; harmless with today's single real
   candidate, but semantically wrong for any future second candidate.
   Fixed here: no schema change, only a different (already-existing)
   column read.

6. APPLY_TODAY now additionally requires: freshness_age_days is known
   and <= MAX_JOB_AGE_DAYS_FOR_APPLY (3, matching this project's
   existing max_job_age_days convention), a usable job_url or
   application_url, and -- for a cross-source duplicate cluster -- only
   the single highest-scoring representative (ties broken by earliest
   first_discovered, then by (source, job_id) for full determinism).
   None of these were enforced by Phase 7's original implementation --
   a stale job, a job with no URL at all, or BOTH sides of a duplicate
   pair could previously appear in APPLY_TODAY simultaneously. Fixed
   here, additively, with no change to eligibility/scoring/freshness/
   dedup logic itself -- only to which already-computed values this
   sheet's membership filter checks.
"""

import json
import os
import sqlite3
import sys
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path


def _parse_timestamp(value):
    """
    Robustly parse a timestamp that may come from either of two
    formats this codebase actually produces:
      - SQLite's own CURRENT_TIMESTAMP: "YYYY-MM-DD HH:MM:SS" (space
        separator, no timezone) -- e.g. jobs.created_at when set by a
        bare `DEFAULT CURRENT_TIMESTAMP` column default.
      - Python's datetime.now(timezone.utc).isoformat(): "YYYY-MM-
        DDTHH:MM:SS.ffffff+00:00" ("T" separator, explicit offset) --
        e.g. jobs.last_updated when set by tracker.py's own `now`.

    Returns a timezone-aware datetime (UTC assumed for the naive
    SQLite form, since every writer in this codebase already uses UTC
    wall-clock time), or None if `value` is empty/unparseable -- never
    raises. This exists because naive STRING comparison between these
    two formats is UNSOUND: "2026-09-20 16:04:48" and
    "2026-09-20T00:00:00" compare as the SPACE character (0x20) sorts
    before "T" (0x54), so a same-day SQLite-formatted timestamp always
    string-compares as "earlier" than an ISO-formatted cutoff on the
    same date, regardless of actual time-of-day -- a real bug this
    module's own Phase 7 SIXTH-step live validation caught (a job
    created at 16:04 the same day as a since="...T00:00:00" cutoff was
    incorrectly excluded from NEW_JOBS).
    """
    if not value:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        if "T" not in text and " " in text:
            text = text.replace(" ", "T", 1)
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from job_ranking import build_ranking_record
from canonical_job import derive_canonical_job
from cross_source_dedup import find_cross_source_duplicate_candidates


# A job is "not yet applied to" -- still a live candidate for APPLY_TODAY.
ACTIVE_STATUSES = ("FOUND", "SCREENING", "SHORTLISTED", "READY_FOR_APPROVAL", "APPROVED")

# A job Saroj has genuinely acted on. Deliberately EXCLUDES "REJECTED"
# and "NOT_QUALIFIED" -- see this module's docstring (Phase 7.2
# section) for the exact evidence-backed reason: both of those values
# are produced exclusively by the score-bucket path
# (scripts/score_job.py), never by any real application-rejection
# workflow. INCLUDES "EMPLOYER_REJECTED" -- unlike the two above, this
# value is unambiguous (no code writes it as a side effect of scoring)
# and, once a real application/approval workflow exists to set it,
# genuinely means "Saroj applied and the employer responded."
APPLIED_LIFECYCLE_STATUSES = (
    "APPLICATION_STARTED",
    "APPLIED",
    "RECRUITER_CONTACTED",
    "SCREENING_CALL",
    "INTERVIEW_1",
    "INTERVIEW_2",
    "FINAL_ROUND",
    "OFFER",
    "EMPLOYER_REJECTED",
    "GHOSTED",
    "WITHDRAWN",
)

# Both values this codebase's score-bucket path has ever written for a
# non-qualifying score: "NOT_QUALIFIED" (current, Phase 7.2 onward) and
# "REJECTED" (legacy -- pre-Phase-7.2 rows only, see this module's
# docstring). Both are treated identically: score-bucket rejection,
# NEVER an implied employer rejection.
SCORE_BUCKET_REJECTION_STATUSES = ("NOT_QUALIFIED", "REJECTED")

QUALIFYING_PRIORITIES = ("A", "B", "C")

# Matches this project's existing max_job_age_days=3 discovery-freshness
# convention (search_profile.py / query_planner.py / naukri_adapter.py).
# APPLY_TODAY (and the "READY_TO_APPLY" Report Status) requires a job's
# RECOMPUTED freshness_age_days to be known and within this bound --
# Phase 7's original implementation did not enforce this at all.
MAX_JOB_AGE_DAYS_FOR_APPLY = 3


@dataclass
class ReportRow:
    ranking: object  # job_ranking.RankingRecord
    status: str
    resume_variant: str
    application_url: str
    first_discovered: str
    last_seen: str
    location_raw: str
    experience_required: str
    # Item 10 fix (export completeness): read directly from the raw job
    # dict, same pattern as location_raw/experience_required above --
    # never invented when the source adapter never populated them.
    work_model: str = ""
    jd_text: str = ""
    # Stamped by generate() AFTER build_report_rows() returns, only for
    # a search-scoped export (job_id_filter/run_ids_filter given) --
    # empty ("") for the candidate-wide report, which has no single
    # search/run to attribute every row to. Never invented.
    search_id: str = ""
    search_run_id: str = ""
    eligible: bool = False
    report_status: str = ""
    previously_seen: object = None  # True/False/None (None = not applicable -- no candidate_job_matches row exists, e.g. non-qualifying jobs)
    duplicate_suppressed: bool = False
    sheets: set = field(default_factory=set)
    # Phase 14, additive: raw jobs.discovery_source/discovery_query/
    # completeness, read directly from the job dict exactly like
    # location_raw/experience_required above -- never derived from
    # `ranking` (job_ranking.RankingRecord has no equivalent fields;
    # not extended for this, since nothing else in the pipeline needs
    # them at that layer -- see this module's docstring note below).
    discovery_source_raw: str = ""
    discovery_query: str = ""
    completeness: object = None


def _row_to_job_dict(row):
    """
    Convert one sqlite3.Row from the `jobs` table into the
    normalize_job()-shaped dict build_ranking_record() expects.
    mandatory_skills/preferred_skills are stored as JSON text in the
    DB; everything else is already the right shape/type.
    """
    job = dict(row)
    for key in ("mandatory_skills", "preferred_skills"):
        raw = job.get(key)
        try:
            job[key] = json.loads(raw) if raw else []
        except (json.JSONDecodeError, TypeError):
            job[key] = []
    return job


def load_jobs(conn, exclude_test_mock=True):
    """
    Read every row from the global `jobs` table. exclude_test_mock=True
    (the default) drops source IN ('TEST', 'MOCK') rows -- the same
    convention generate_daily_report.py already uses.
    """
    conn.row_factory = sqlite3.Row
    query = "SELECT * FROM jobs"
    if exclude_test_mock:
        query += " WHERE source NOT IN ('TEST', 'MOCK')"
    rows = conn.execute(query).fetchall()
    return [_row_to_job_dict(row) for row in rows]


def load_candidate_job_matches(conn, candidate_id):
    """
    Read every candidate_job_matches row for this candidate, keyed by
    job_id. This table only ever has a row for a job that PASSED
    eligibility (search_worker.upsert_candidate_job_match() is only
    called from the eligible branch of process_queue_item()) -- so a
    lookup miss here always means "never eligible for this candidate,"
    never "not yet processed."

    created_at on this table is protected the same way jobs.created_at
    is (confirmed by reading upsert_candidate_job_match()'s ON CONFLICT
    UPDATE SET clause -- created_at is not in it) -- so it is a genuine,
    per-candidate "first time this candidate matched this job" signal,
    unlike jobs.created_at (global, not candidate-scoped).
    """
    conn.row_factory = sqlite3.Row
    rows = conn.execute(
        "SELECT job_id, created_at, updated_at, candidate_status, resume_variant FROM candidate_job_matches WHERE candidate_id = ?",
        (candidate_id,),
    ).fetchall()
    return {row["job_id"]: dict(row) for row in rows}


def _identity_key(job):
    return (job.get("source"), job.get("job_id"))


def _build_duplicate_clusters(duplicate_candidates):
    """
    Union-find over (source, job_id) identities connected by ANY
    reported duplicate-candidate pair -- handles a job matched to more
    than one other source transitively, not just pairwise. Returns
    {identity_key: cluster_root_key} for every identity that appears in
    at least one pair.
    """
    parent = {}

    def find(x):
        parent.setdefault(x, x)
        root = x
        while parent[root] != root:
            root = parent[root]
        while parent[x] != root:
            parent[x], x = root, parent[x]
        return root

    def union(a, b):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[ra] = rb

    for dup in duplicate_candidates:
        a = (dup.job_a.source, dup.job_a.source_job_id)
        b = (dup.job_b.source, dup.job_b.source_job_id)
        union(a, b)

    return {key: find(key) for key in parent}


def _pick_duplicate_representatives(clusters, rankings_by_key, jobs_by_key):
    """
    For each duplicate cluster, pick exactly ONE representative identity
    to remain eligible for APPLY_TODAY: highest score first (None
    treated as lowest), tie-broken by earliest first_discovered
    (jobs.created_at), tie-broken by (source, job_id) for full
    determinism. Returns the set of representative identity keys --
    every OTHER identity in a cluster is "suppressed" from APPLY_TODAY
    (still visible everywhere else: ALL_MATCHING_JOBS, DUPLICATES,
    APPLICATION_TRACKER).
    """
    members_by_root = {}
    for key, root in clusters.items():
        members_by_root.setdefault(root, []).append(key)

    representatives = set()
    for members in members_by_root.values():
        def _sort_key(key):
            ranking = rankings_by_key.get(key)
            score = ranking.score if ranking and ranking.score is not None else -1
            created = _parse_timestamp((jobs_by_key.get(key) or {}).get("created_at"))
            created_sort = created.isoformat() if created else "9999"
            return (-score, created_sort, key)

        best = min(members, key=_sort_key)
        representatives.add(best)

    return representatives


def _has_usable_url(job, ranking):
    return bool((job.get("application_url") or "").strip() or (ranking.job_url or "").strip())


def _compute_report_status(ranking, status, in_applied_lifecycle, is_apply_today_worthy, duplicate_suppressed):
    """
    One of: SCORE_REJECTED_NOT_QUALIFIED / ELIGIBLE_NOT_APPLIED /
    READY_TO_APPLY / APPLIED / <literal applied-lifecycle status,
    including EMPLOYER_REJECTED/WITHDRAWN/GHOSTED> / DUPLICATE /
    INELIGIBLE / AMBIGUOUS_REJECTED_STATUS_SEE_LIMITATION. Pure
    derivation from already-computed values -- no new eligibility/
    scoring logic, no schema change.

    Phase 7.2 (data/reports/phase7_2_automation_status_audit.md):
    scripts/score_job.py now writes "NOT_QUALIFIED", never "REJECTED",
    for a score-bucket rejection -- but existing production rows
    written before this change may still say "REJECTED" (confirmed:
    3 such rows exist in production today, all source='TEST'). Both
    values (SCORE_BUCKET_REJECTION_STATUSES) are handled IDENTICALLY
    here, forever -- this is a forward-only vocabulary change, not a
    migration, and old rows must remain correctly classified without
    ever being reinterpreted as an employer rejection.

    "EMPLOYER_REJECTED" is now a real, unambiguous, reserved status
    (no code writes it yet, but nothing here needs to change once
    something does) -- it flows through the generic
    `in_applied_lifecycle` branch below like any other real lifecycle
    status, returned verbatim.
    """
    if duplicate_suppressed:
        return "DUPLICATE"

    if not ranking.eligible:
        return "INELIGIBLE"

    qualifies = ranking.priority in QUALIFYING_PRIORITIES

    if status in SCORE_BUCKET_REJECTION_STATUSES:
        # This codebase can only currently produce a score-bucket
        # rejection status via score_job()'s own bucket assignment (at
        # insert) or as a stale leftover from an earlier lower score
        # (status is protected from re-overwrite on re-scoring). It
        # can NEVER currently represent a genuine employer rejection --
        # neither the current "NOT_QUALIFIED" nor the legacy "REJECTED".
        if qualifies:
            # Reachable only if a job was first scored REJECT (status
            # auto-set to NOT_QUALIFIED/REJECTED), then later re-scored
            # higher -- priority updates on re-score, status does not.
            # A stale artifact, not a real signal either way --
            # flagged, not guessed at, and excluded from
            # APPLY_TODAY/READY_TO_APPLY.
            return "AMBIGUOUS_REJECTED_STATUS_SEE_LIMITATION"
        return "SCORE_REJECTED_NOT_QUALIFIED"

    if in_applied_lifecycle:
        return status  # APPLIED / RECRUITER_CONTACTED / INTERVIEW_1 / ... / EMPLOYER_REJECTED / WITHDRAWN / GHOSTED -- all unambiguous

    if not qualifies:
        return "SCORE_REJECTED_NOT_QUALIFIED"

    return "READY_TO_APPLY" if is_apply_today_worthy else "ELIGIBLE_NOT_APPLIED"


def build_report_rows(jobs, candidate_profile, since=None, candidate_job_matches=None):
    """
    Pure, offline. jobs: list of dicts from load_jobs(). candidate_profile:
    the flat legacy-shaped dict candidate_profile.to_legacy_matching_profile()
    produces. candidate_job_matches: dict from load_candidate_job_matches()
    (job_id -> row), or None/empty if unavailable -- NEW_JOBS/`Previously
    Seen` are then never populated for any row (an explicit, no-silent-
    default choice, not a guess). since: an ISO datetime string cutoff
    compared against candidate_job_matches.created_at (per-candidate --
    see Phase 7.1 docstring section, point 5), or None ("NEW_JOBS never
    populated").

    Returns (report_rows, duplicate_candidates).
    """
    candidate_job_matches = candidate_job_matches or {}
    jobs_by_key = {_identity_key(job): job for job in jobs}

    canonical_jobs = [derive_canonical_job(job) for job in jobs]
    duplicate_candidates = find_cross_source_duplicate_candidates(canonical_jobs)

    duplicates_by_identity = {}
    for dup in duplicate_candidates:
        for cjob in (dup.job_a, dup.job_b):
            duplicates_by_identity.setdefault((cjob.source, cjob.source_job_id), []).append(dup)

    # Pass 1: compute every ranking first -- duplicate-representative
    # selection needs every cluster member's score up front.
    rankings_by_key = {}
    for job in jobs:
        key = _identity_key(job)
        rankings_by_key[key] = build_ranking_record(
            job, candidate_profile, duplicate_candidates=duplicates_by_identity.get(key, [])
        )

    clusters = _build_duplicate_clusters(duplicate_candidates)
    representatives = _pick_duplicate_representatives(clusters, rankings_by_key, jobs_by_key)

    since_dt = _parse_timestamp(since) if since else None

    report_rows = []
    for job in jobs:
        key = _identity_key(job)
        ranking = rankings_by_key[key]

        # THIS candidate's own lifecycle status (candidate_job_matches.
        # candidate_status -- set on discovery from score_job()'s
        # bucket, and moved forward from there by a human via
        # PATCH /api/candidates/{id}/jobs/{id}/status) is authoritative
        # when it exists. jobs.status is a single-candidate-era GLOBAL
        # column that gets overwritten by whichever candidate most
        # recently had this worker score this job (see search_worker.
        # py's _upsert_job_scored() docstring) -- it is only a
        # fallback for a job this candidate was never matched to at all
        # (ineligible -- candidate_job_matches never has a row for
        # that, per load_candidate_job_matches()'s own docstring).
        match_row = candidate_job_matches.get(job.get("job_id"))
        status = (match_row or {}).get("candidate_status") or str(job.get("status") or "FOUND")
        created_at = str(job.get("created_at") or "")
        last_seen = str(job.get("last_updated") or job.get("updated_at") or created_at)

        qualifies = ranking.priority in QUALIFYING_PRIORITIES
        is_applied = status in APPLIED_LIFECYCLE_STATUSES

        previously_seen = None
        is_new = False
        if match_row is not None and since_dt is not None:
            match_created_dt = _parse_timestamp(match_row.get("created_at"))
            if match_created_dt is not None:
                is_new = match_created_dt >= since_dt
                previously_seen = not is_new

        in_cluster = key in clusters
        duplicate_suppressed = in_cluster and key not in representatives

        is_fresh_enough = (
            ranking.freshness_age_days is not None
            and ranking.freshness_age_days <= MAX_JOB_AGE_DAYS_FOR_APPLY
        )
        has_url = _has_usable_url(job, ranking)
        is_apply_today_worthy = (
            qualifies
            and status in ACTIVE_STATUSES
            and is_fresh_enough
            and has_url
            and not duplicate_suppressed
        )

        report_status = _compute_report_status(
            ranking, status, is_applied, is_apply_today_worthy, duplicate_suppressed
        )

        # THIS candidate's own selected resume (candidate_job_matches.
        # resume_variant, populated by resume_variant_selector.py) is
        # authoritative when set -- jobs.resume_variant is the same
        # single-candidate-era GLOBAL leftover as jobs.status above,
        # and stays permanently empty in practice (see
        # migrate_v7_candidate_resume_variant.py).
        resume_variant = (match_row or {}).get("resume_variant") or str(job.get("resume_variant") or "")

        row = ReportRow(
            ranking=ranking,
            status=status,
            resume_variant=resume_variant,
            application_url=str(job.get("application_url") or ""),
            first_discovered=created_at,
            last_seen=last_seen,
            location_raw=str(job.get("location") or ""),
            experience_required=str(job.get("experience_required") or ""),
            work_model=str(job.get("work_model") or ""),
            jd_text=str(job.get("jd_text") or ""),
            eligible=ranking.eligible,
            report_status=report_status,
            previously_seen=previously_seen,
            duplicate_suppressed=duplicate_suppressed,
            discovery_source_raw=str(job.get("discovery_source") or ""),
            discovery_query=str(job.get("discovery_query") or ""),
            completeness=job.get("completeness"),
        )

        sheets = {"APPLICATION_TRACKER"}

        if is_applied:
            sheets.add("ALREADY_APPLIED")
        elif qualifies:
            sheets.add("ALL_MATCHING_JOBS")
            if is_apply_today_worthy:
                sheets.add("APPLY_TODAY")
            if is_new:
                sheets.add("NEW_JOBS")
        else:
            sheets.add("REJECTED_EXCLUDED")

        if ranking.duplicate_candidates:
            sheets.add("DUPLICATES")

        row.sheets = sheets
        report_rows.append(row)

    return report_rows, duplicate_candidates


def build_source_health(conn, candidate_id=None, run_ids_filter=None):
    """
    SOURCE_HEALTH sheet source: the persisted search_runs table -- the
    only place per-run source/query/error/block counts survive after a
    process exits (WorkItemResult itself is in-memory only).

    run_ids_filter (item 10/11 fix -- run-scoped export): an optional
    set/container of search_run_ids. When given, restricts this to
    exactly those runs -- e.g. a per-search report must not let
    build_run_summary()'s "most recent run" picture come from a
    DIFFERENT search's more-recent run for the same candidate. None
    (the default) preserves the original candidate-wide behavior.
    """
    conn.row_factory = sqlite3.Row
    query = """
        SELECT search_run_id, candidate_id, source, role, location, status,
               started_at, completed_at, queries_total, queries_completed,
               jobs_discovered, jobs_deduplicated, jobs_experience_excluded,
               jobs_eligible, jobs_scored, jobs_ready, errors_count,
               blocked_queries, error_message, created_at
        FROM search_runs
    """
    clauses = []
    params = []
    if candidate_id is not None:
        clauses.append("candidate_id = ?")
        params.append(candidate_id)
    if run_ids_filter is not None:
        run_ids_list = list(run_ids_filter)
        if not run_ids_list:
            return []
        clauses.append(f"search_run_id IN ({','.join('?' for _ in run_ids_list)})")
        params.extend(run_ids_list)
    if clauses:
        query += " WHERE " + " AND ".join(clauses)
    query += " ORDER BY created_at DESC"

    return [dict(row) for row in conn.execute(query, params).fetchall()]


def build_run_summary(report_rows, source_health):
    """
    Answers Phase 7's REPORT QUALITY ACCEPTANCE TEST question 12 ("how
    many jobs were discovered and how many survived each filtering
    stage?") for the MOST RECENT run, plus current-state distributions
    across every row in scope (all runs' accumulated jobs).
    """
    priority_distribution = Counter(
        row.ranking.priority or "INELIGIBLE" for row in report_rows
    )
    freshness_distribution = Counter(row.ranking.freshness for row in report_rows)
    report_status_distribution = Counter(row.report_status for row in report_rows)

    scored_values = [row.ranking.score for row in report_rows if row.ranking.score is not None]

    latest_run = source_health[0] if source_health else None

    return {
        "latest_run": latest_run,
        "total_jobs_in_scope": len(report_rows),
        "eligible_count": sum(1 for row in report_rows if row.ranking.eligible),
        "ineligible_count": sum(1 for row in report_rows if not row.ranking.eligible),
        "qualifying_count": sum(1 for row in report_rows if row.ranking.priority in QUALIFYING_PRIORITIES),
        "priority_distribution": dict(priority_distribution),
        "freshness_distribution": dict(freshness_distribution),
        "report_status_distribution": dict(report_status_distribution),
        "score_min": min(scored_values) if scored_values else None,
        "score_max": max(scored_values) if scored_values else None,
        "duplicate_candidate_pairs": sum(
            1 for row in report_rows if row.ranking.duplicate_candidates
        )
        // 2,  # each pair counted from both sides
        "apply_today_count": sum(1 for row in report_rows if "APPLY_TODAY" in row.sheets),
        "new_jobs_count": sum(1 for row in report_rows if "NEW_JOBS" in row.sheets),
        "already_applied_count": sum(1 for row in report_rows if "ALREADY_APPLIED" in row.sheets),
        "rejected_excluded_count": sum(1 for row in report_rows if "REJECTED_EXCLUDED" in row.sheets),
    }


# ---------------------------------------------------------------------
# .xlsx workbook rendering
# ---------------------------------------------------------------------

def _previously_seen_display(r):
    if r.previously_seen is None:
        return "N/A"
    return "Yes" if r.previously_seen else "No"


def _discovered_via_display(r):
    """"DIRECT" for a blank discovery_source, or one equal to the job's
    own source (discover_local.normalize_job()'s own default for every
    caller that never set discovery_source explicitly -- every direct
    adapter, and every job written before Phase 14's columns existed).
    Anything else (e.g. "SEARCH_PROVIDER:SERPER", "MANUAL_IMPORT") is
    shown verbatim -- never re-derived or guessed here."""
    raw = r.discovery_source_raw
    if not raw or raw == r.ranking.source:
        return "DIRECT"
    return raw


_JOB_SHEET_COLUMNS = [
    ("Priority", lambda r: r.ranking.priority or ""),
    ("Score", lambda r: r.ranking.score if r.ranking.score is not None else ""),
    ("Title", lambda r: r.ranking.title),
    ("Company", lambda r: r.ranking.company),
    ("Location", lambda r: r.location_raw),
    ("Source", lambda r: r.ranking.source),
    ("Job URL", lambda r: r.ranking.job_url),
    ("Application URL", lambda r: r.application_url or r.ranking.job_url),
    ("Experience Required", lambda r: r.experience_required),
    ("Freshness", lambda r: r.ranking.freshness),
    ("Freshness (days)", lambda r: r.ranking.freshness_age_days if r.ranking.freshness_age_days is not None else ""),
    ("Eligible", lambda r: "Yes" if r.eligible else "No"),
    ("Report Status", lambda r: r.report_status),
    ("Matched Skills / Strong Points", lambda r: ", ".join((r.ranking.score_explanation or {}).get("strong_matches", []))),
    ("Gaps", lambda r: ", ".join((r.ranking.score_explanation or {}).get("gaps", []))),
    ("Exclusion / Gap Reason", lambda r: "; ".join(r.ranking.eligibility_reasons)),
    ("Application Status", lambda r: r.status),
    ("Previously Seen", _previously_seen_display),
    ("Duplicate (Suppressed from Apply Today)", lambda r: "Yes" if r.duplicate_suppressed else "No"),
    ("First Discovered", lambda r: r.first_discovered),
    ("Last Seen", lambda r: r.last_seen),
    ("Resume Variant", lambda r: r.resume_variant),
    # Phase 14 additions -- appended at the end so no existing column's
    # position/letter shifts for anyone reading this workbook by
    # position or an existing formula referencing a column letter.
    ("Discovered_Via", _discovered_via_display),
    ("Discovery_Query", lambda r: r.discovery_query),
    ("Completeness", lambda r: r.completeness if r.completeness is not None else ""),
    # Informational only (experience_eligibility.detect_requirement_type)
    # -- never affects Eligible/Priority/Report Status above. Appended
    # at the end, same convention as the Phase 14 columns, so no
    # existing column letter shifts.
    ("Requirement Type", lambda r: r.ranking.requirement_type),
    # Item 10 fix (export completeness): a stable, unique job
    # identifier was previously absent from every export entirely --
    # appended at the end, same convention as every other addition
    # above, so no existing column letter shifts for anyone with an
    # existing formula/reference into this workbook.
    ("Job ID", lambda r: r.ranking.job_id),
    ("Work Model", lambda r: r.work_model),
    ("JD Text", lambda r: r.jd_text),
    # Export completeness fix: blank for the candidate-wide report
    # (no single search/run to attribute every row to); populated only
    # for a search-scoped export, via generate()'s search_id/
    # search_run_id params -- same appended-at-the-end convention.
    ("Search ID", lambda r: r.search_id),
    ("Search Run ID", lambda r: r.search_run_id),
]

_URL_COLUMNS = {"Job URL", "Application URL"}


def _sort_key_apply_priority(row):
    priority_rank = {"A": 0, "B": 1, "C": 2}.get(row.ranking.priority, 9)
    score = row.ranking.score if row.ranking.score is not None else -1
    return (priority_rank, -score)


def _write_job_sheet(wb, name, rows, sort_key=_sort_key_apply_priority):
    ws = wb.create_sheet(name)
    headers = [label for label, _ in _JOB_SHEET_COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _BOLD
    ws.freeze_panes = "A2"

    for row in sorted(rows, key=sort_key):
        values = [getter(row) for _, getter in _JOB_SHEET_COLUMNS]
        ws.append(values)
        written = ws[ws.max_row]
        for header, cell in zip(headers, written):
            if header in _URL_COLUMNS and cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    _autosize(ws)
    return ws


def _write_duplicates_sheet(wb, duplicate_candidates):
    ws = wb.create_sheet("DUPLICATES")
    headers = [
        "Title", "Company", "Confidence", "Signals", "Reason",
        "Source A", "Job URL A", "Source B", "Job URL B",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _BOLD
    ws.freeze_panes = "A2"

    for dup in duplicate_candidates:
        ws.append([
            dup.job_a.normalized_title,
            dup.job_a.normalized_company,
            dup.confidence,
            ", ".join(dup.signals),
            dup.reason,
            dup.job_a.source,
            dup.job_a.job_url,
            dup.job_b.source,
            dup.job_b.job_url,
        ])
        written = ws[ws.max_row]
        for header, cell in zip(headers, written):
            if header in ("Job URL A", "Job URL B") and cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    _autosize(ws)
    return ws


def _write_source_health_sheet(wb, source_health):
    ws = wb.create_sheet("SOURCE_HEALTH")
    if not source_health:
        ws.append(["No search_runs rows found for this candidate/DB."])
        return ws

    headers = list(source_health[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _BOLD
    ws.freeze_panes = "A2"
    for entry in source_health:
        ws.append([entry.get(h, "") for h in headers])

    _autosize(ws)
    return ws


def _write_run_summary_sheet(wb, summary):
    ws = wb.create_sheet("RUN_SUMMARY")
    ws.append(["Metric", "Value"])
    for cell in ws[1]:
        cell.font = _BOLD

    def _flatten(prefix, value):
        if isinstance(value, dict):
            for k, v in value.items():
                _flatten(f"{prefix}.{k}" if prefix else str(k), v)
        else:
            ws.append([prefix, value if value is not None else ""])

    for key, value in summary.items():
        _flatten(key, value)

    _autosize(ws)
    return ws


def _autosize(ws, max_width=60):
    for column_cells in ws.columns:
        length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=8)
        col_letter = column_cells[0].column_letter
        ws.column_dimensions[col_letter].width = min(max_width, max(10, length + 2))


def generate_workbook(report_rows, duplicate_candidates, source_health, summary, out_path):
    """
    Writes the 9-sheet workbook. Returns the Path written.

    Phase 7.2 (data/reports/phase7_2_automation_status_audit.md, PART B
    requirement #13): ATOMIC write -- the workbook is built and saved
    to a temporary file in the SAME directory as `out_path` (so the
    final rename is on the same filesystem, hence atomic on POSIX),
    and only `Path.replace()`d onto `out_path` after `wb.save()`
    completes without error. If anything fails partway (a bad row, a
    disk-full condition, an interrupted process), `out_path` -- and
    therefore any PREVIOUS successful workbook already sitting there --
    is never touched; the incomplete temp file is cleaned up instead.
    This is what lets a daily scheduler run safely: a failed run can
    never corrupt or truncate the last good report.
    """
    import openpyxl
    from openpyxl.styles import Font

    global _BOLD
    _BOLD = Font(bold=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    _write_job_sheet(wb, "APPLY_TODAY", [r for r in report_rows if "APPLY_TODAY" in r.sheets])
    _write_job_sheet(wb, "ALL_MATCHING_JOBS", [r for r in report_rows if "ALL_MATCHING_JOBS" in r.sheets])
    _write_job_sheet(wb, "NEW_JOBS", [r for r in report_rows if "NEW_JOBS" in r.sheets])
    _write_job_sheet(wb, "ALREADY_APPLIED", [r for r in report_rows if "ALREADY_APPLIED" in r.sheets], sort_key=lambda r: r.last_seen or "")
    _write_job_sheet(wb, "REJECTED_EXCLUDED", [r for r in report_rows if "REJECTED_EXCLUDED" in r.sheets], sort_key=lambda r: -(r.ranking.score or -1))
    _write_duplicates_sheet(wb, duplicate_candidates)
    _write_job_sheet(wb, "APPLICATION_TRACKER", report_rows, sort_key=lambda r: r.last_seen or "")
    _write_source_health_sheet(wb, source_health)
    _write_run_summary_sheet(wb, summary)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    tmp_path = out_path.with_name(f".{out_path.name}.tmp-{os.getpid()}")
    try:
        wb.save(tmp_path)
        tmp_path.replace(out_path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise
    return out_path


# ---------------------------------------------------------------------
# Top-level orchestration
# ---------------------------------------------------------------------

def generate(db_path, candidate_id, out_path, since=None, exclude_test_mock=True, planner_metrics=None, job_id_filter=None, run_ids_filter=None, resume_overrides_by_job=None, search_id=None, search_run_id=None):
    """
    db_path: path to a SQLite DB with the v2 schema (candidates,
    candidate_search_profile, jobs, search_runs, ...).
    candidate_id: whose confirmed profile to score against.
    out_path: where to write the .xlsx workbook.
    since: ISO datetime string cutoff for NEW_JOBS, or None (see
    build_report_rows()'s docstring).

    job_id_filter (item 10/11 fix -- run-scoped export): an optional
    set/container of job_ids. When given, the workbook is restricted to
    exactly those jobs -- e.g. api/main.py's per-search report endpoint
    passes api/results_store.py's OWN _job_ids_for_search() output, the
    SAME canonical scoping function the results JSON API uses, so the
    export can never show a different job count than the UI for the
    same search/run. None (the default) preserves the original
    behavior exactly: every job in scope for this candidate, matching
    the existing candidate-wide report endpoint. Cross-source
    deduplication (build_report_rows()) always runs across the FULL,
    unfiltered job set first -- filtering to job_id_filter happens
    only after, exactly mirroring get_results_for_saved_search()'s own
    order of operations -- so a duplicate's cluster-mate outside this
    filter is still correctly accounted for, never mis-flagged.
    planner_metrics (Phase 14.6, optional): a flat/nested dict of
    daily-search-planner metrics (planned/skipped/executed queries,
    provider requests saved, cooldown/budget skips, new vs. known jobs,
    etc. -- see daily_search_planner.py). When given, merged into the
    summary dict under the "planner" key before the RUN_SUMMARY sheet
    is written. _write_run_summary_sheet() already flattens `summary`
    generically, so this adds rows to the EXISTING RUN_SUMMARY sheet --
    no new worksheet, no change to any of the other 8 sheets or to any
    existing summary key. None (the default) leaves the summary exactly
    as every prior phase produced it.

    resume_overrides_by_job (item 5/10 fix, same class of bug already
    fixed for the JSON results API in api/results_store.py's
    _apply_scoped_resume_overrides()): an optional {job_id: {"resume_id":
    ..., "resume_variant": ...}} dict. When given, overrides the
    candidate-wide "latest snapshot" resume_variant/resume_id for the
    matching job_ids -- otherwise a per-search export could show a
    DIFFERENT search's more recently-run resume, exactly the bug this
    project's own audit already found and fixed for the JSON API. None
    (the default) preserves the original candidate-wide behavior.

    search_id / search_run_id (export completeness fix): the saved
    search / search_run this export is scoped to (via job_id_filter/
    run_ids_filter above). Constant across every row in a search-scoped
    export, so stamped onto each ReportRow after build_report_rows()
    returns rather than looked up per-row. None (the default) leaves
    both columns blank -- the candidate-wide report has no single
    search/run to attribute every row to, and this must never be
    guessed.

    Returns a dict summary (also embedded as the RUN_SUMMARY sheet).
    """
    from search_submission import load_candidate_confirmed_profile
    from candidate_profile import to_legacy_matching_profile

    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys = ON")
    try:
        candidate_profile = load_candidate_confirmed_profile(conn, candidate_id)
        legacy_profile = to_legacy_matching_profile(candidate_profile)

        jobs = load_jobs(conn, exclude_test_mock=exclude_test_mock)
        matches = load_candidate_job_matches(conn, candidate_id)
        if resume_overrides_by_job:
            for job_id, override in resume_overrides_by_job.items():
                matches.setdefault(job_id, {})
                matches[job_id]["resume_id"] = override.get("resume_id")
                matches[job_id]["resume_variant"] = override.get("resume_variant")
        report_rows, duplicate_candidates = build_report_rows(
            jobs, legacy_profile, since=since, candidate_job_matches=matches
        )
        if job_id_filter is not None:
            report_rows = [row for row in report_rows if row.ranking.job_id in job_id_filter]
        if search_id is not None or search_run_id is not None:
            for row in report_rows:
                row.search_id = str(search_id) if search_id is not None else ""
                row.search_run_id = str(search_run_id) if search_run_id is not None else ""
        source_health = build_source_health(conn, candidate_id=candidate_id, run_ids_filter=run_ids_filter)
        summary = build_run_summary(report_rows, source_health)
        if planner_metrics is not None:
            summary["planner"] = planner_metrics

        workbook_path = generate_workbook(report_rows, duplicate_candidates, source_health, summary, out_path)
    finally:
        conn.close()

    return {
        "workbook_path": str(workbook_path),
        "summary": summary,
        "row_count": len(report_rows),
    }


def _default_since():
    today = datetime.now(timezone.utc).date().isoformat()
    return f"{today}T00:00:00"


def main():
    import argparse

    parser = argparse.ArgumentParser(description="Generate the Phase 7 multi-sheet daily job report workbook.")
    parser.add_argument("--db", required=True, help="Path to the SQLite DB (v2 schema).")
    parser.add_argument("--candidate-id", required=True)
    parser.add_argument("--out", required=True, help="Output .xlsx path.")
    parser.add_argument("--since", default=None, help="ISO datetime cutoff for NEW_JOBS (default: start of today, UTC).")
    parser.add_argument("--include-test-mock", action="store_true", help="Include source IN ('TEST','MOCK') rows.")
    args = parser.parse_args()

    since = args.since if args.since is not None else _default_since()

    result = generate(
        args.db, args.candidate_id, args.out, since=since, exclude_test_mock=not args.include_test_mock
    )

    print("RUN REPORT GENERATED")
    print("=====================")
    print(f"Workbook       : {result['workbook_path']}")
    print(f"Jobs in scope  : {result['row_count']}")
    summary = result["summary"]
    print(f"APPLY_TODAY    : {summary['apply_today_count']}")
    print(f"NEW_JOBS       : {summary['new_jobs_count']}")
    print(f"ALREADY_APPLIED: {summary['already_applied_count']}")
    print(f"REJECTED       : {summary['rejected_excluded_count']}")
    print(f"Priority dist. : {summary['priority_distribution']}")
    print(f"Freshness dist.: {summary['freshness_distribution']}")


if __name__ == "__main__":
    main()
