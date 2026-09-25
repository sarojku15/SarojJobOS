#!/usr/bin/env python3

"""
Phase 8 PART F -- multi-source aggregation / cross-source deduplication
test (data/reports/phase8_multisource_validation.md).

cross_source_dedup.py and canonical_job.py are SOURCE-AGNOSTIC -- they
operate on already-normalized job dicts regardless of which adapter
produced them, and regardless of whether that adapter's
AdapterStatus is ENABLED (that flag only gates LIVE FETCHING via
source_registry, never processing of an already-obtained job dict).
This lets Part F be tested fully offline, with zero live requests to
either Naukri or Hirist: synthetic job dicts shaped exactly like each
source's real normalize_job() output are used instead. LinkedIn is
excluded per Part E's conclusion (LinkedIn is not enabled, and no
authorized access mechanism exists to source real or representative
data from it).

Never opens data/applications/jobos.db. Never makes a network call.
"""

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import generate_run_report as grr

_RICH_JD = (
    "Senior Site Reliability Engineer / DevOps role. Own production reliability "
    "across AWS and Azure. Run Kubernetes (EKS/AKS) clusters, manage "
    "infrastructure as code with Terraform and Ansible, build CI/CD pipelines "
    "with Jenkins and GitHub Actions, and drive incident management, "
    "SLO/SLI/error budget practices, and root cause analysis. Observability "
    "via Prometheus and Grafana is core."
)

CANDIDATE_PROFILE = {
    "candidate": {"experience_years": 11},
    "target_locations": ["Bengaluru"],
    "cloud": ["AWS", "Azure"],
    "kubernetes": ["Kubernetes", "EKS", "AKS"],
    "iac_and_automation": ["Terraform", "Ansible"],
    "cicd_and_devops": ["Jenkins", "GitHub Actions"],
    "observability": ["Prometheus", "Grafana"],
    "sre": [],
    "programming_and_scripting": [],
    "tools": [],
}


def _job(source, job_id, company, title, url_suffix, jd_text=_RICH_JD, location="Bengaluru"):
    return {
        "source": source, "job_id": job_id, "company": company, "title": title,
        "location": location, "work_model": "Hybrid",
        "job_url": f"https://{source.lower()}.example.com/jobs/{url_suffix}",
        "application_url": f"https://{source.lower()}.example.com/apply/{url_suffix}",
        "posted_date": "2026-09-20", "jd_text": jd_text, "experience_required": "8-12 years",
        "mandatory_skills": [], "preferred_skills": [],
    }


# 1. Naukri-only job -- unique, no duplicate anywhere
JOB_NAUKRI_ONLY = _job("NAUKRI", "N-ONLY-001", "Alpha Systems", "Senior Site Reliability Engineer", "alpha-1")

# 2. Hirist-only job -- unique, no duplicate anywhere
JOB_HIRIST_ONLY = _job("HIRIST", "H-ONLY-001", "Beta Cloud", "Senior Site Reliability Engineer", "beta-1")

# 3/4. Same job on Naukri + Hirist (identical title/company/location/JD)
JOB_SAME_NAUKRI = _job("NAUKRI", "N-SAME-002", "Gamma Corp", "Senior Site Reliability Engineer", "gamma-naukri")
JOB_SAME_HIRIST = _job("HIRIST", "H-SAME-002", "Gamma Corp", "Senior Site Reliability Engineer", "gamma-hirist")

# 5. Same job, but with a tracking-parameter-decorated URL on one side --
# cross_source_dedup.compare_pair() never compares job_url at all (only
# normalized_title/normalized_company/canonical_location/description),
# so this must dedupe identically to case 3/4 regardless of URL shape.
JOB_TRACKING_NAUKRI = _job("NAUKRI", "N-TRACK-003", "Delta Infra", "Senior Site Reliability Engineer", "delta-naukri")
JOB_TRACKING_HIRIST = dict(_job("HIRIST", "H-TRACK-003", "Delta Infra", "Senior Site Reliability Engineer", "delta-hirist"))
JOB_TRACKING_HIRIST["job_url"] = JOB_TRACKING_HIRIST["job_url"] + "?utm_source=newsletter&utm_campaign=weekly&ref=abc123"

# 6. Similar title, DIFFERENT company -- must NOT be flagged as duplicate
JOB_SIMILAR_DIFFERENT_COMPANY_A = _job("NAUKRI", "N-DIFFCO-004", "Epsilon Tech", "Senior Site Reliability Engineer", "epsilon-1")
JOB_SIMILAR_DIFFERENT_COMPANY_B = _job("HIRIST", "H-DIFFCO-004", "Zeta Cloud", "Senior Site Reliability Engineer", "zeta-1")

# 7. Same company/title text, but genuinely different jobs (different
# location) -- must NOT be flagged as duplicate (canonical_location
# mismatch overrides an otherwise-matching title+company pair)
JOB_SAME_TITLE_DIFFERENT_LOC_A = _job("NAUKRI", "N-SAMETXT-005", "Theta Corp", "Senior Site Reliability Engineer", "theta-blr", location="Bengaluru")
JOB_SAME_TITLE_DIFFERENT_LOC_B = _job("HIRIST", "H-SAMETXT-005", "Theta Corp", "Senior Site Reliability Engineer", "theta-hyd", location="Hyderabad")

ALL_JOBS = [
    JOB_NAUKRI_ONLY, JOB_HIRIST_ONLY,
    JOB_SAME_NAUKRI, JOB_SAME_HIRIST,
    JOB_TRACKING_NAUKRI, JOB_TRACKING_HIRIST,
    JOB_SIMILAR_DIFFERENT_COMPANY_A, JOB_SIMILAR_DIFFERENT_COMPANY_B,
    JOB_SAME_TITLE_DIFFERENT_LOC_A, JOB_SAME_TITLE_DIFFERENT_LOC_B,
]


def _fail(failures, message):
    failures.append(message)
    print(f"FAIL: {message}")


def main():
    failures = []
    print("PHASE 8 PART F -- MULTI-SOURCE AGGREGATION / CROSS-SOURCE DEDUP TEST (offline)")
    print("=================================================================================")

    report_rows, duplicate_candidates = grr.build_report_rows(ALL_JOBS, CANDIDATE_PROFILE, since=None)
    rows_by_id = {(r.ranking.source, r.ranking.job_id): r for r in report_rows}

    def sheets_of(source, job_id):
        return rows_by_id[(source, job_id)].sheets

    # 1. Naukri-only job: no duplicate
    if "DUPLICATES" in sheets_of("NAUKRI", "N-ONLY-001"):
        _fail(failures, "1. Naukri-only job incorrectly flagged as duplicate")
    else:
        print("PASS: 1. Naukri-only job correctly has no duplicate")

    # 2. Hirist-only job: no duplicate
    if "DUPLICATES" in sheets_of("HIRIST", "H-ONLY-001"):
        _fail(failures, "2. Hirist-only job incorrectly flagged as duplicate")
    else:
        print("PASS: 2. Hirist-only job correctly has no duplicate (dedup logic is source-agnostic -- works identically for a NOT_ENABLED source's data shape)")

    # 3/4. Same job on Naukri + Hirist: both flagged
    if "DUPLICATES" not in sheets_of("NAUKRI", "N-SAME-002") or "DUPLICATES" not in sheets_of("HIRIST", "H-SAME-002"):
        _fail(failures, f"3/4. expected both sides flagged DUPLICATES: naukri={sheets_of('NAUKRI', 'N-SAME-002')} hirist={sheets_of('HIRIST', 'H-SAME-002')}")
    else:
        print("PASS: 3/4. the same job posted on both Naukri and Hirist is correctly detected as a cross-source duplicate candidate")

    # 5. Tracking-URL variant: still deduped identically (URL never compared)
    if "DUPLICATES" not in sheets_of("NAUKRI", "N-TRACK-003") or "DUPLICATES" not in sheets_of("HIRIST", "H-TRACK-003"):
        _fail(failures, f"5. expected both sides flagged DUPLICATES despite tracking-parameter URL differences: naukri={sheets_of('NAUKRI', 'N-TRACK-003')} hirist={sheets_of('HIRIST', 'H-TRACK-003')}")
    else:
        print("PASS: 5. a tracking-parameter-decorated URL does not prevent duplicate detection (cross_source_dedup.compare_pair() never compares job_url at all -- confirmed by code and by this test)")

    # 6. Similar title, different company: NOT flagged
    if "DUPLICATES" in sheets_of("NAUKRI", "N-DIFFCO-004") or "DUPLICATES" in sheets_of("HIRIST", "H-DIFFCO-004"):
        _fail(failures, "6. genuinely different jobs (different company) were incorrectly flagged as duplicates")
    else:
        print("PASS: 6. same title, different company -> correctly NOT flagged as duplicate")

    # 7. Same company/title text, different location: NOT flagged
    if "DUPLICATES" in sheets_of("NAUKRI", "N-SAMETXT-005") or "DUPLICATES" in sheets_of("HIRIST", "H-SAMETXT-005"):
        _fail(failures, "7. same company/title but genuinely different jobs (different location) were incorrectly flagged as duplicates")
    else:
        print("PASS: 7. same company/title text but different location -> correctly NOT flagged as duplicate (canonical_location mismatch overrides an otherwise-matching pair)")

    # Overall reconciliation: exactly 2 duplicate pairs found (3/4 and 5),
    # never more (no false positives elsewhere), never fewer.
    if len(duplicate_candidates) != 2:
        _fail(failures, f"reconciliation: expected exactly 2 duplicate-candidate pairs (case 3/4 and case 5), got {len(duplicate_candidates)}")
    else:
        print(f"PASS: reconciliation -- exactly 2 duplicate-candidate pairs found across {len(ALL_JOBS)} synthetic jobs, matching the 2 genuinely-duplicate cases exactly (no false positives, no false negatives)")

    # No unsafe double-counting: ALL_MATCHING_JOBS / APPLY_TODAY must not
    # count both sides of a genuine duplicate as independently actionable
    # -- only the representative (higher-scoring / earliest / (source,job_id)
    # tie-broken) side should remain in APPLY_TODAY.
    apply_today_ids = {(r.ranking.source, r.ranking.job_id) for r in report_rows if "APPLY_TODAY" in r.sheets}
    duplicate_pair_keys = [
        {("NAUKRI", "N-SAME-002"), ("HIRIST", "H-SAME-002")},
        {("NAUKRI", "N-TRACK-003"), ("HIRIST", "H-TRACK-003")},
    ]
    for pair in duplicate_pair_keys:
        overlap = pair & apply_today_ids
        if len(overlap) > 1:
            _fail(failures, f"unsafe double-counting: both sides of a duplicate pair appear in APPLY_TODAY: {overlap}")
    else:
        print("PASS: no unsafe double-counting -- at most one side of each genuine duplicate pair appears in APPLY_TODAY")

    # === PART G: generate the actual 9-sheet workbook from this same
    # multi-source dataset and verify structure/content ===
    import tempfile
    import openpyxl

    source_health = []  # no search_runs in this purely offline, no-DB scenario
    summary = grr.build_run_summary(report_rows, source_health)
    out_path = Path(tempfile.mkdtemp(prefix="jobos_phase8_partg_")) / "multisource_report.xlsx"
    grr.generate_workbook(report_rows, duplicate_candidates, source_health, summary, out_path)

    wb = openpyxl.load_workbook(out_path)
    expected_sheets = {
        "APPLY_TODAY", "ALL_MATCHING_JOBS", "NEW_JOBS", "ALREADY_APPLIED",
        "REJECTED_EXCLUDED", "DUPLICATES", "APPLICATION_TRACKER", "SOURCE_HEALTH", "RUN_SUMMARY",
    }
    if set(wb.sheetnames) != expected_sheets:
        _fail(failures, f"PART G: workbook sheet names mismatch: got {set(wb.sheetnames)}")
    else:
        print(f"PASS: PART G -> multi-source workbook has exactly the 9 required sheets: {sorted(wb.sheetnames)}")

    apply_today_ws = wb["APPLY_TODAY"]
    headers = [c.value for c in apply_today_ws[1]]
    required_columns = ["Priority", "Score", "Title", "Company", "Location", "Source", "Job URL",
                         "Freshness", "Eligible", "Report Status", "Application Status",
                         "Duplicate (Suppressed from Apply Today)"]
    missing_columns = [c for c in required_columns if c not in headers]
    if missing_columns:
        _fail(failures, f"PART G: APPLY_TODAY sheet missing required column(s): {missing_columns}")
    else:
        print(f"PASS: PART G -> APPLY_TODAY exposes all required columns: source, title, company, location, score, priority, freshness, URL, eligibility, duplicate state, application status")

    # Confirm the DUPLICATES sheet lists both genuine duplicate pairs
    duplicates_ws = wb["DUPLICATES"]
    duplicate_row_count = duplicates_ws.max_row - 1
    if duplicate_row_count != 2:
        _fail(failures, f"PART G: DUPLICATES sheet expected 2 rows, got {duplicate_row_count}")
    else:
        print("PASS: PART G -> DUPLICATES sheet lists exactly 2 rows, matching the 2 genuine duplicate pairs")

    # Confirm APPLY_TODAY never contains a duplicate-suppressed, stale,
    # unusable-URL, already-applied, or non-qualifying job (all synthetic
    # jobs here are fresh/eligible/qualifying/have URLs by construction,
    # so this specifically re-confirms the duplicate-suppression rule
    # using REAL sheet content, not just in-memory row.sheets checks).
    apply_today_company_col_idx = headers.index("Company") + 1
    apply_today_companies = [row[apply_today_company_col_idx - 1] for row in apply_today_ws.iter_rows(min_row=2, values_only=True)]
    if apply_today_companies.count("Gamma Corp") > 1 or apply_today_companies.count("Delta Infra") > 1:
        _fail(failures, f"PART G: APPLY_TODAY contains BOTH sides of a duplicate pair: {apply_today_companies}")
    else:
        print("PASS: PART G -> APPLY_TODAY never contains both sides of any duplicate pair (verified from actual workbook content)")

    if failures:
        print("\nFAILURES:")
        for f in failures:
            print(f"  {f}")
        sys.exit(1)

    print("\nAll Phase 8 PART F/G multi-source dedup + reporting tests passed.")


if __name__ == "__main__":
    main()
