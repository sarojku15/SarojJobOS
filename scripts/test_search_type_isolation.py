#!/usr/bin/env python3

"""
Regression test for a real bug found live on this project's own dev
DB: the normal /searches page and Dashboard had no way to distinguish
a genuine user search from a development/verification search, and the
Dashboard's "Jobs Found"/"Matching" aggregates were not even scoped by
candidate -- they read the ENTIRE global `jobs` table (1555 rows on
the real dev DB, shared across every candidate/search/live-
verification run this project has ever made), not the viewing
candidate's own activity.

Fix: saved_searches.search_type (USER/TEST/SYSTEM, migrate_v10_
search_type.py), defaulting to USER for every real create-search call;
api/results_store.py's get_dashboard_summary() now scopes to this
candidate's own candidate_job_matches, minus any job tracked only
through one of their own TEST-type searches.

Covers:
  1. A TEST search is excluded from the normal /searches listing.
  2. A TEST search's jobs never inflate the Dashboard's aggregates.
  3. A USER search remains fully visible on /searches and contributes
     normally to the Dashboard.
  4. The per-search Excel export is scoped to exactly the jobs the
     results API shows for THAT search/run -- never a different,
     larger candidate-wide count (the second real bug found: the
     per-search report endpoint used to return the SAME whole-
     candidate report as /api/candidates/{id}/report).
  5. Export columns exist and are non-fabricated.

Fully offline (fake adapters, no real network call). Never opens
data/applications/jobos.db.
"""

import hashlib
import io
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = ROOT / "scripts"
API_DIR = ROOT / "api"
sys.path.insert(0, str(SCRIPTS_DIR))
sys.path.insert(0, str(API_DIR))

PRODUCTION_DB = ROOT / "data" / "applications" / "jobos.db"


def _sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


production_before = _sha(PRODUCTION_DB)

import source_registry
from source_adapter import JobSourceAdapter, AdapterHealth, AdapterStatus, BlockReason

passed = 0
failed = 0


def check(condition, message):
    global passed, failed
    if condition:
        passed += 1
        print(f"PASS: {message}")
    else:
        failed += 1
        print(f"FAIL: {message}")


class _IsolationAdapterUser(JobSourceAdapter):
    name = "FAKE_ISOLATION_SOURCE"
    status = AdapterStatus.ENABLED

    def health_check(self):
        return AdapterHealth(source=self.name, reachable=True, block_reason=BlockReason.NONE)

    def search(self, query):
        return [
            {
                "source": self.name, "job_id": "FAKE-ISOLATION-USER-JOB", "company": "UserSearchCo",
                "title": query.role, "location": query.location, "work_model": "Remote",
                "job_url": "https://example.com/jobs/isolation-user", "application_url": "https://example.com/apply/isolation-user",
                "posted_date": "2026-09-20", "jd_text": "Senior Cloud Engineer role. kubernetes.",
                "experience_required": "3+ years", "mandatory_skills": [], "preferred_skills": [],
            }
        ]


class _IsolationAdapterTest(JobSourceAdapter):
    name = "FAKE_ISOLATION_SOURCE"
    status = AdapterStatus.ENABLED

    def health_check(self):
        return AdapterHealth(source=self.name, reachable=True, block_reason=BlockReason.NONE)

    def search(self, query):
        return [
            {
                "source": self.name, "job_id": "FAKE-ISOLATION-TEST-JOB", "company": "TestSearchCo",
                "title": query.role, "location": query.location, "work_model": "Remote",
                "job_url": "https://example.com/jobs/isolation-test", "application_url": "https://example.com/apply/isolation-test",
                "posted_date": "2026-09-20", "jd_text": "Senior Cloud Engineer role. kubernetes terraform.",
                "experience_required": "3+ years", "mandatory_skills": [], "preferred_skills": [],
            }
        ]


source_registry.ADAPTERS["FAKE_ISOLATION_SOURCE"] = _IsolationAdapterUser
_original_list_sources = source_registry.list_sources
source_registry.list_sources = lambda: ["FAKE_ISOLATION_SOURCE"]

try:
    tmp_dir = Path(tempfile.mkdtemp(prefix="jobos_test_search_type_"))
    tmp_db = tmp_dir / "jobos_test.db"

    import init_dev_db
    init_dev_db.init_dev_db(tmp_db)

    import db as db_mod
    db_mod.DEV_DB = tmp_db

    from fastapi.testclient import TestClient
    import main as api_main
    api_main.db_mod.DEV_DB = tmp_db

    client = TestClient(api_main.app)

    resp = client.post("/api/candidates", json={"name": "Isolation Test Candidate"})
    candidate_id = resp.json()["candidate_id"]
    client.put(
        f"/api/candidates/{candidate_id}/profile",
        json={
            "current_title": "Senior Cloud Engineer", "total_experience_years": 6,
            "skills": {"containers_orchestration": [{"name": "Kubernetes"}]},
            "job_preferences": {"target_roles": ["Senior Cloud Engineer"], "target_locations": ["Remote"]},
        },
    )
    client.post(f"/api/candidates/{candidate_id}/profile/confirm")

    # --- create a USER search (the default -- no search_type given) ---
    resp = client.post(
        f"/api/candidates/{candidate_id}/searches",
        json={
            "name": "Real user search", "target_roles": ["Senior Cloud Engineer"], "target_locations": ["Remote"],
            "minimum_match_score": 0, "sources": ["FAKE_ISOLATION_SOURCE"],
        },
    )
    user_search = resp.json()
    user_search_id = user_search["saved_search_id"]
    check(user_search.get("search_type", "USER") == "USER" or "search_type" not in user_search, f"setup: real create-search call defaults to USER (no search_type in payload), got field={user_search.get('search_type')!r}")

    # Confirm directly in the DB too (in case the response doesn't echo it).
    conn = db_mod.get_conn()
    row = conn.execute("SELECT search_type FROM saved_searches WHERE saved_search_id = ?", (user_search_id,)).fetchone()
    conn.close()
    check(row[0] == "USER", f"1/3 setup: DB confirms the real create-search call stored search_type=USER by default, got {row[0]!r}")

    # --- create a TEST search explicitly ---
    resp = client.post(
        f"/api/candidates/{candidate_id}/searches",
        json={
            "name": "Automated verification search", "target_roles": ["Senior Cloud Engineer"], "target_locations": ["Remote"],
            "minimum_match_score": 0, "sources": ["FAKE_ISOLATION_SOURCE"], "search_type": "TEST",
        },
    )
    test_search = resp.json()
    test_search_id = test_search["saved_search_id"]
    check(test_search.get("search_type") == "TEST", f"setup: explicit search_type=TEST is honored, got {test_search.get('search_type')!r}")

    # Invalid search_type is rejected.
    resp = client.post(
        f"/api/candidates/{candidate_id}/searches",
        json={
            "name": "Bad type", "target_roles": ["X"], "target_locations": ["Y"],
            "minimum_match_score": 0, "search_type": "NOT_A_REAL_TYPE",
        },
    )
    check(resp.status_code == 400, f"an invalid search_type is rejected (400), got {resp.status_code}")

    # --- 1/3. /searches listing shows USER, never TEST ---
    resp = client.get(f"/api/candidates/{candidate_id}/searches")
    listed_names = {s["name"] for s in resp.json()["searches"]}
    check("Real user search" in listed_names, "3. the USER search is visible on the normal /searches listing")
    check("Automated verification search" not in listed_names, "1. the TEST search is EXCLUDED from the normal /searches listing")

    # --- run both searches ---
    def run_and_wait(search_id):
        resp = client.post(f"/api/searches/{search_id}/run?candidate_id={candidate_id}")
        run_id = resp.json()["run_id"]
        deadline = time.time() + 20
        status = None
        while time.time() < deadline:
            r = client.get(f"/api/runs/{run_id}?candidate_id={candidate_id}")
            status = r.json()["status"]
            if status in ("COMPLETED", "PARTIAL", "FAILED", "BLOCKED"):
                break
            time.sleep(0.2)
        return run_id, status

    source_registry.ADAPTERS["FAKE_ISOLATION_SOURCE"] = _IsolationAdapterUser
    user_run_id, user_status = run_and_wait(user_search_id)
    check(user_status == "COMPLETED", f"setup: USER search run completes, got {user_status}")

    source_registry.ADAPTERS["FAKE_ISOLATION_SOURCE"] = _IsolationAdapterTest
    test_run_id, test_status = run_and_wait(test_search_id)
    check(test_status == "COMPLETED", f"setup: TEST search run completes, got {test_status}")

    # --- 2. Dashboard aggregates reflect ONLY the USER search's job,
    #     never the TEST search's job ---
    resp = client.get(f"/api/candidates/{candidate_id}/dashboard")
    dash = resp.json()
    dash_search_names = {s["name"] for s in dash["searches"]}
    check("Real user search" in dash_search_names, "3. Dashboard's own per-search panel shows the USER search")
    check("Automated verification search" not in dash_search_names, "2. Dashboard's own per-search panel excludes the TEST search")

    summary = dash["summary"]
    check(summary["total_jobs_in_scope"] == 1, f"2. Dashboard 'Jobs Found' (total_jobs_in_scope) counts ONLY the USER search's 1 job, NOT the TEST search's job too, got {summary['total_jobs_in_scope']}")
    check(summary["eligible_count"] == 1, f"2. Dashboard 'Matching' (eligible_count) also excludes the TEST search's job, got {summary['eligible_count']}")

    # --- 4. per-search Excel export is scoped to exactly the jobs the
    #     results API shows for that search ---
    resp = client.get(f"/api/searches/{user_search_id}/results?candidate_id={candidate_id}")
    ui_results = resp.json()["results"]
    ui_job_ids = {r["job_id"] for r in ui_results}

    resp = client.get(f"/api/searches/{user_search_id}/report?candidate_id={candidate_id}")
    check(resp.status_code == 200, f"export: 200, got {resp.status_code}")
    check(resp.headers.get("content-type", "").startswith("application/vnd.openxmlformats"), f"export: real xlsx content-type, got {resp.headers.get('content-type')}")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        sheet_names = wb.sheetnames
        # A job appears in exactly one of ALL_MATCHING_JOBS (qualified,
        # priority A/B/C) / REJECTED_EXCLUDED (hard-eligible but not
        # qualified) / ALREADY_APPLIED, by this project's own existing,
        # intentional sheet-splitting design (build_report_rows()'s
        # `sheets` set) -- so "the export contains the same jobs as the
        # UI" means the UNION across every job-bearing sheet, not any
        # one sheet alone (the UI's single results list shows every
        # hard-eligible job regardless of qualification).
        job_bearing_sheets = ["ALL_MATCHING_JOBS", "REJECTED_EXCLUDED", "ALREADY_APPLIED"]
        exported_job_ids = set()
        header_row = None
        for name in job_bearing_sheets:
            if name not in sheet_names:
                continue
            ws = wb[name]
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            if header_row is None:
                header_row = header
            job_id_col = next((idx for idx, h in enumerate(header) if h and "job" in str(h).lower() and "id" in str(h).lower()), None)
            if job_id_col is None:
                continue
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[job_id_col]:
                    exported_job_ids.add(row[job_id_col])

        check(exported_job_ids == ui_job_ids, f"4/14/15. export (union of ALL_MATCHING_JOBS/REJECTED_EXCLUDED/ALREADY_APPLIED) contains EXACTLY the same job_ids as the UI's results for this search/run -- UI={ui_job_ids}, export={exported_job_ids}")

        expected_columns = {"title", "company", "location", "source", "score", "priority"}
        found_columns_lower = {str(h).lower() for h in header_row if h} if header_row else set()
        missing = [c for c in expected_columns if not any(c in fc for fc in found_columns_lower)]
        check(not missing, f"16. export sheet has the expected core columns, missing: {missing} (actual header: {header_row})")
    except ImportError:
        check(False, "openpyxl not available to verify export contents (test environment gap, not a product defect)")

finally:
    source_registry.list_sources = _original_list_sources

production_after = _sha(PRODUCTION_DB)
check(production_before == production_after, f"production DB byte-identical before/after this test run (sha256 before={production_before}, after={production_after})")

print()
print(f"{passed} passed, {failed} failed")
if failed:
    sys.exit(1)
